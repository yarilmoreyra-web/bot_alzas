"""Datos del día para toda la watchlist, hasta la hora de la consulta:
cierre anterior, apertura, precio actual, volumen, máximo y mínimo.

Envía a Telegram SOLO un archivo Excel (datos_dia.xlsx) con todos los tickers.

Usa la lista de tickers y las funciones de consulta.py (deben estar juntos
en la misma carpeta / repositorio)."""
import os
import time
import logging
from datetime import datetime

import requests
from longbridge.openapi import QuoteContext
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from consulta import (
    WATCHLIST, TZ_NY, TZ_LIMA, TELEGRAM_BOT_TOKEN, TELEGRAM_CHAT_ID,
    enviar_telegram, crear_config, sesion_actual,
)

log = logging.getLogger("datos_dia")

NOMBRE_ARCHIVO = "datos_dia.xlsx"
FUENTE = "Arial"

ETIQUETAS = {
    None: "regular",
    "pre_market_quote": "pre-market",
    "post_market_quote": "after-hours",
    "overnight_quote": "overnight",
}


def a_float(x):
    try:
        return float(x) if x else 0.0
    except (TypeError, ValueError):
        return 0.0


def obtener_quotes(ctx, simbolos, tam_lote=100):
    """Cotizaciones por lotes; si un lote falla, reintenta símbolo por símbolo."""
    resultado = []
    for i in range(0, len(simbolos), tam_lote):
        lote = simbolos[i:i + tam_lote]
        try:
            resultado.extend(ctx.quote(lote))
        except Exception as e:
            log.warning("Falló el lote (%s). Reintentando uno a uno...", e)
            for s in lote:
                try:
                    resultado.extend(ctx.quote([s]))
                except Exception as e2:
                    log.warning("Símbolo omitido %s: %s", s, e2)
                time.sleep(0.12)
    return resultado


def leer_datos(q, sesion):
    """Datos de un ticker. Vol, máximo y mínimo son de la sesión en curso.
    La apertura es la de la sesión regular (no existe aún en pre-market)."""
    fuente = q if sesion is None else getattr(q, sesion, None)
    if not fuente:
        return None
    ahora = a_float(fuente.last_done)
    if ahora <= 0:
        return None
    ant = a_float(q.prev_close)
    return {
        "simbolo": q.symbol,
        "ant": ant,
        "aper": 0.0 if sesion == "pre_market_quote" else a_float(q.open),
        "ahora": ahora,
        "pct": (ahora - ant) / ant * 100 if ant > 0 else None,
        "vol": int(fuente.volume or 0),
        "max": a_float(getattr(fuente, "high", 0)),
        "min": a_float(getattr(fuente, "low", 0)),
    }


def crear_excel(filas, sin_datos, etiqueta, ahora_ny, hora_txt, ruta):
    """Hoja 'Datos' con una fila por ticker y hoja 'Info' con el contexto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    ws.append(["Símbolo", "Cierre anterior", "Apertura", "Precio actual", "Variación %",
               "Volumen", "Máximo", "Mínimo", "Sesión", "Hora consulta (ET)"])
    for d in filas:
        ws.append([
            d["simbolo"],
            d["ant"] or None,
            d["aper"] or None,
            d["ahora"],
            d["pct"] / 100 if d["pct"] is not None else None,   # fracción: 0.026 = 2.60%
            d["vol"],
            d["max"] or None,
            d["min"] or None,
            etiqueta,
            ahora_ny.strftime("%Y-%m-%d %H:%M"),
        ])

    # Formato
    relleno = PatternFill("solid", start_color="1F3864")
    for c in ws[1]:
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    formatos = {2: "$#,##0.00", 3: "$#,##0.00", 4: "$#,##0.00", 5: "0.00%",
                6: "#,##0", 7: "$#,##0.00", 8: "$#,##0.00"}
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=FUENTE)
            if c.column in formatos:
                c.number_format = formatos[c.column]

    for i, ancho in enumerate([11, 15, 12, 14, 13, 16, 12, 12, 13, 19], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # Hoja con el contexto de la consulta
    info = wb.create_sheet("Info")
    filas_info = [
        ("Consulta", hora_txt),
        ("Sesión en curso", etiqueta),
        ("Tickers con datos", len(filas)),
        ("Tickers sin datos", ", ".join(sin_datos) if sin_datos else "ninguno"),
        ("", ""),
        ("Cierre anterior", "Cierre de la sesión regular anterior (según Longbridge)."),
        ("Apertura", "Precio de apertura de la sesión regular. Vacío en pre-market."),
        ("Precio actual", "Último precio de la sesión en curso."),
        ("Variación %", "(Precio actual − Cierre anterior) / Cierre anterior."),
        ("Volumen, Máximo, Mínimo", "De la sesión en curso, hasta la hora de la consulta."),
        ("Orden", "De mayor a menor variación vs el cierre anterior."),
    ]
    for etiqueta_fila, valor in filas_info:
        info.append([etiqueta_fila, valor])
    for fila in info.iter_rows():
        fila[0].font = Font(name=FUENTE, bold=True)
        fila[1].font = Font(name=FUENTE)
        fila[1].alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 70

    wb.save(ruta)


def enviar_documento(ruta, caption):
    """Devuelve True si Telegram aceptó el archivo."""
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        log.warning("Faltan tokens de Telegram.")
        return False
    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"
    try:
        with open(ruta, "rb") as f:
            r = requests.post(url, data={"chat_id": TELEGRAM_CHAT_ID, "caption": caption},
                              files={"document": f}, timeout=60)
        if not r.ok:
            log.error("Telegram respondió %s: %s", r.status_code, r.text)
        return r.ok
    except (requests.RequestException, OSError) as e:
        log.error("No pude enviar el archivo: %s", e)
        return False


def main():
    ahora_ny = datetime.now(TZ_NY)
    hora_txt = f"{ahora_ny.strftime('%H:%M')} ET / {ahora_ny.astimezone(TZ_LIMA).strftime('%H:%M')} Lima"
    sesion = sesion_actual(ahora_ny)
    etiqueta = ETIQUETAS[sesion]
    log.info("Datos del día, sesión %s, %d tickers", etiqueta, len(WATCHLIST))

    try:
        ctx = QuoteContext(crear_config())
    except Exception as e:
        log.exception("No se pudo conectar con Longbridge")
        enviar_telegram(f"⚠️ No pude conectar con Longbridge: {e}", parse_mode=None)
        raise

    filas = []
    for q in obtener_quotes(ctx, WATCHLIST):
        try:
            d = leer_datos(q, sesion)
        except (TypeError, ValueError, AttributeError) as e:
            log.warning("Dato inválido en %s: %s", getattr(q, "symbol", "?"), e)
            continue
        if d:
            filas.append(d)

    if not filas:
        enviar_telegram(f"⚠️ Sin datos de la sesión {etiqueta} a las {hora_txt}.", parse_mode=None)
        return

    filas.sort(key=lambda d: d["pct"] if d["pct"] is not None else -1e9, reverse=True)
    con_datos = {d["simbolo"] for d in filas}
    sin_datos = [s for s in WATCHLIST if s not in con_datos]

    crear_excel(filas, sin_datos, etiqueta, ahora_ny, hora_txt, NOMBRE_ARCHIVO)

    caption = f"Datos del día · {len(filas)} tickers · {hora_txt} · sesión {etiqueta}"
    if sin_datos:
        caption += f" · sin datos: {len(sin_datos)} (ver hoja Info)"
    if enviar_documento(NOMBRE_ARCHIVO, caption):
        log.info("Archivo enviado (%d tickers, %d sin datos).", len(filas), len(sin_datos))
    else:
        enviar_telegram("⚠️ Generé el archivo pero Telegram no lo aceptó. Revisa el registro en GitHub.",
                        parse_mode=None)


if __name__ == "__main__":
    main()
